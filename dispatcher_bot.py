import os
import io
import re
import json
import logging
import mimetypes
import html
import asyncio
import random
import httpx
from google import genai
from google.genai import types
from telegram import Update
from telegram.ext import Application, CommandHandler, MessageHandler, ConversationHandler, filters

TELEGRAM_TOKEN = os.environ.get("TELEGRAM_TOKEN")
GEMINI_KEY = os.environ.get("GEMINI_KEY")
GOOGLE_MAPS_KEY = os.environ.get("GOOGLE_MAPS_KEY")
BOT_PASSWORD = os.environ.get("BOT_PASSWORD")  # Fly.io Secrets
AUTH_FILE = os.environ.get("AUTH_FILE", "authorized_users.json")

WAITING_LOCATION = 1

logging.basicConfig(format="%(asctime)s | %(levelname)s | %(message)s", level=logging.INFO)
logger = logging.getLogger(__name__)

client = genai.Client(api_key=GEMINI_KEY)
MODEL_NAME = "gemini-3.1-flash-lite"
model = MODEL_NAME
chat_model = MODEL_NAME

MAX_RETRIES = 5
BASE_DELAY = 5  # seconds

# --- Authentication -------------------------------------------------------
def _load_authorized():
    try:
        with open(AUTH_FILE, "r") as f:
            return set(json.load(f))
    except Exception:
        return set()


def _save_authorized():
    try:
        with open(AUTH_FILE, "w") as f:
            json.dump(sorted(AUTHORIZED_USERS), f)
    except Exception:
        logger.warning("Could not save authorized users file")


AUTHORIZED_USERS = _load_authorized()


def is_authorized(user_id):
    return user_id in AUTHORIZED_USERS


async def request_password(update):
    await update.message.reply_text("🔒 Bu bot himoyalangan.\nIltimos, parolni kiriting:")


async def handle_password_attempt(update, context):
    if not update.message:
        return
    text = (update.message.text or "").strip()
    if not text:
        await request_password(update)
        return
    if BOT_PASSWORD and text == BOT_PASSWORD:
        AUTHORIZED_USERS.add(update.effective_user.id)
        _save_authorized()
        # parol chatda qolmasligi uchun xabarni o'chiramiz
        try:
            await update.message.delete()
        except Exception:
            pass
        await context.bot.send_message(
            chat_id=update.effective_chat.id,
            text="✅ Parol to'g'ri! Endi botdan foydalanishingiz mumkin.\n\n/start buyrug'ini yuboring.",
        )
    else:
        await update.message.reply_text("❌ Parol noto'g'ri. Qayta urinib ko'ring:")


def require_auth(handler_func):
    async def wrapped(update, context, *args, **kwargs):
        user = update.effective_user
        if user is None or not is_authorized(user.id):
            await handle_password_attempt(update, context)
            return ConversationHandler.END
        return await handler_func(update, context, *args, **kwargs)
    return wrapped

# ---------------------------------------------------------------------------


def _is_rate_limit_error(exc):
    msg = str(exc)
    return (
        "429" in msg
        or "503" in msg
        or "UNAVAILABLE" in msg
        or "quota" in msg.lower()
        or ("rate" in msg.lower() and "limit" in msg.lower())
    )


def _extract_retry_delay(exc):
    match = re.search(r"retry_delay\s*\{\s*seconds:\s*(\d+)", str(exc))
    if match:
        return int(match.group(1))
    match = re.search(r"retry in ([\d.]+)s", str(exc), re.IGNORECASE)
    if match:
        return float(match.group(1))
    return None


async def generate_with_retry(model_name, contents):
    last_exc = None
    for attempt in range(MAX_RETRIES):
        try:
            return await asyncio.to_thread(
                client.models.generate_content, model=model_name, contents=contents
            )
        except Exception as e:
            last_exc = e
            if not _is_rate_limit_error(e):
                raise
            suggested = _extract_retry_delay(e)
            delay = suggested if suggested is not None else BASE_DELAY * (2 ** attempt)
            delay = delay + random.uniform(0, 1)
            logger.warning(
                "Gemini rate-limited/unavailable (attempt %s/%s), retrying in %.1fs",
                attempt + 1, MAX_RETRIES, delay,
            )
            await asyncio.sleep(delay)
    raise last_exc


def parse_json_response(text):
    raw = (text or "").strip()
    raw = re.sub(r"^```(?:json)?\s*", "", raw)
    raw = re.sub(r"```$", "", raw).strip()
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        start, end = raw.find("{"), raw.rfind("}")
        if start != -1 and end > start:
            return json.loads(raw[start:end + 1])
        raise


EXTRACT_PROMPT = """You are a freight dispatcher assistant.
Analyze this load/rate confirmation and return ONLY a valid JSON object, no explanation, no markdown, no backticks.
JSON format:
{
  "broker": "broker company name only",
  "load_number": "load number",
  "equipment": "equipment type",
  "weight": "total weight with unit",
  "rate": "total rate e.g. $2,500.00",
  "rate_per_mile": "rate per mile e.g. $3.38/mile or empty string",
  "pickups": [
    {
      "number": 1,
      "facility": "facility name or empty string",
      "address_line1": "street address only e.g. 14901 N Beach St",
      "address_line2": "city, state, zip e.g. Fort Worth, TX, 76177",
      "date": "MM/DD/YYYY",
      "time": "time or ASAP",
      "instruction": "LIVE LOAD or DROP or empty",
      "commodity": "commodity description or empty",
      "vrid": "VRID number if shown in document, else empty string",
      "pu_number": "PU number/reference if shown in document, else empty string",
      "bol": "BOL number if shown in document, else empty string",
      "appt": "appointment number/confirmation number if shown in document, else empty string",
      "po": "PO number if shown in document, else empty string"
    }
  ],
  "deliveries": [
    {
      "number": 1,
      "facility": "facility name or empty string",
      "address_line1": "street address only e.g. 124 Davis St",
      "address_line2": "city, state, zip e.g. Portland, TN, 37148",
      "date": "MM/DD/YYYY",
      "time": "time or ASAP",
      "instruction": "LIVE UNLOAD or DROP or empty",
      "commodity": "commodity description or empty",
      "vrid": "VRID number if shown in document, else empty string",
      "pu_number": "PU number/reference if shown in document, else empty string",
      "bol": "BOL number if shown in document, else empty string",
      "appt": "appointment number/confirmation number if shown in document, else empty string",
      "po": "PO number if shown in document, else empty string"
    }
  ]
}
Rules:
- Include ALL pickup and delivery stops
- facility: location name/code
- address_line1: street only
- address_line2: city, state, zip
- date: always format as MM/DD/YYYY with a 4-digit year
- time: copy exactly as shown in the document (e.g. "0900" or "0900-1100")
- instruction/commodity: empty string if not shown in the document
- vrid/pu_number/bol/appt/po: only fill in if explicitly shown in the document for that specific stop; otherwise return an empty string. Do not guess or invent values.
- Return ONLY the JSON, nothing else"""

EDIT_PROMPT = """You are a freight dispatcher assistant helping edit load data.
Below is the CURRENT load data as JSON, followed by an instruction from the dispatcher
asking to change something about it (e.g. "change DO time to 3:00 PM", "PU 1 sanasini ertaga qo'y").

Decide:
1. If the dispatcher's message is clearly asking to CHANGE/EDIT/UPDATE/CORRECT something in the load data
   (times, dates, addresses, facility names, broker, load number, equipment, weight, rate, instructions,
   commodity, VRID/PU#/BOL#/Appt#/PO#, etc.), apply the change(s) to the JSON and return ONLY the updated
   JSON object, with the EXACT same structure/fields as the input, no explanation, no markdown, no backticks.
2. If the message is NOT an edit request (it's a question, comment, or unrelated chat), return ONLY this
   exact JSON: {{"__not_an_edit__": true}}

Only change the field(s) the dispatcher actually asked about. Keep every other field exactly as it was.
If the dispatcher refers to "DO" they mean a delivery stop; "PU" means a pickup stop. If they don't specify
a stop number and there is only one pickup or one delivery, apply it to that one. If there are multiple and
it's ambiguous which one, make your best reasonable guess based on the message.

CURRENT LOAD DATA:
{load_data}

DISPATCHER MESSAGE:
{user_text}

Return ONLY the JSON object described above, nothing else."""


def get_mime(filename):
    mime, _ = mimetypes.guess_type(filename)
    return mime or "application/octet-stream"


def _esc(value):
    return html.escape(str(value), quote=False)


def _ref_lines(stop):
    fields = [
        ("VRID#", stop.get("vrid", "")),
        ("PU#:", stop.get("pu_number", "")),
        ("BOL#", stop.get("bol", "")),
        ("Appt#", stop.get("appt", "")),
        ("PO#", stop.get("po", "")),
    ]
    return [f"❕{label} {_esc(value)}" for label, value in fields if value]


def _stop_block(lines, stop, header):
    lines.append("")
    lines.append(header)
    if stop.get("facility"):
        lines.append(_esc(stop["facility"]))
    if stop.get("address_line1"):
        lines.append(_esc(stop["address_line1"]))
    if stop.get("address_line2"):
        lines.append(_esc(stop["address_line2"]))
    lines.append(f"📅Date: {_esc(stop.get('date', ''))}")
    lines.append(f"🕔Time : {_esc(stop.get('time', ''))}")
    code_lines = []
    if stop.get("instruction"):
        code_lines.append(f"🚛 Instruction:{_esc(stop['instruction'])}")
    if stop.get("commodity"):
        code_lines.append(f"📤Commodity: {_esc(stop['commodity'])}")
    code_lines.extend(_ref_lines(stop))
    if code_lines:
        lines.append(f"<code>{chr(10).join(code_lines)}</code>")


def build_message(load_data, empty_miles, loaded_miles):
    lines = []
    lines.append(f"📌Broker:  {_esc(load_data.get('broker', ''))}")
    lines.append("Al Amin Express Inc")
    lines.append(f"Load:    {_esc(load_data.get('load_number', ''))}")

    for pu in load_data.get("pickups", []):
        _stop_block(lines, pu, f"🟢PU {pu.get('number', '')} :")

    for do_ in load_data.get("deliveries", []):
        _stop_block(lines, do_, f"🔴DO {do_.get('number', '')}:")

    lines.append("")
    lines.append(f"Empty :  {empty_miles} mile")
    lines.append(f"Loaded :  {loaded_miles} mile")

    lines.append("")
    lines.append("❌MUST SEND TRAILER PICTURES, TRAILER REGISTRATION PAPER, TRAILER REGISTRATION PAPER, TO THE GROUP AND WAIT FOR THE GOOD TO GO CONFIRMATION BY THE ONLY DISPATCHER/UPDATER.")
    lines.append("DO NOT DEPART WITHOUT CONFIRMATION!!!")
    lines.append("")
    lines.append("❌LATE PICKUP $500 DEDUCTION!!!")
    lines.append("")
    lines.append("❌LATE DELIVERY $700 DEDUCTION!!!")
    lines.append("")
    lines.append("❌MUST USE AMAZON RELAY")

    return "\n".join(lines)


async def get_distance_miles(origin, destination):
    url = "https://maps.googleapis.com/maps/api/distancematrix/json"
    params = {
        "origins": origin,
        "destinations": destination,
        "units": "imperial",
        "key": GOOGLE_MAPS_KEY,
    }
    try:
        async with httpx.AsyncClient(timeout=10) as http:
            r = await http.get(url, params=params)
        data = r.json()
        element = data["rows"][0]["elements"][0]
        if element["status"] != "OK":
            return 0
        return round(element["distance"]["value"] / 1609.344)
    except Exception:
        logger.exception("Distance Matrix error")
        return 0


async def calculate_miles(current_location, load_data):
    all_stops = []
    for pu in load_data.get("pickups", []):
        all_stops.append(pu.get("address_line1", "") + ", " + pu.get("address_line2", ""))
    for do_ in load_data.get("deliveries", []):
        all_stops.append(do_.get("address_line1", "") + ", " + do_.get("address_line2", ""))
    if not all_stops:
        return 0, 0
    # barcha so'rovlar bir vaqtda (tezroq)
    tasks = [get_distance_miles(current_location, all_stops[0])]
    tasks += [get_distance_miles(all_stops[i], all_stops[i + 1]) for i in range(len(all_stops) - 1)]
    results = await asyncio.gather(*tasks)
    return results[0], sum(results[1:])


def _read_docx(file_bytes):
    import docx as _docx
    doc = _docx.Document(io.BytesIO(file_bytes))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append("\t".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def _read_xlsx(file_bytes):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    rows = []
    for sheet in wb.worksheets:
        rows.append("=== " + sheet.title + " ===")
        for row in sheet.iter_rows(values_only=True):
            rows.append("\t".join(str(c) if c is not None else "" for c in row))
    return "\n".join(rows)


async def extract_load_data(file_bytes, filename, caption=""):
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    mime = get_mime(filename)

    if mime.startswith("image/"):
        part = types.Part.from_bytes(data=file_bytes, mime_type=mime)
        response = await generate_with_retry(model, [EXTRACT_PROMPT, part])
    elif mime == "application/pdf" or ext == "pdf":
        part = types.Part.from_bytes(data=file_bytes, mime_type="application/pdf")
        response = await generate_with_retry(model, [EXTRACT_PROMPT, part])
    elif ext == "docx":
        try:
            text = _read_docx(file_bytes)
        except Exception as e:
            text = "[Could not parse: " + str(e) + "]"
        response = await generate_with_retry(model, EXTRACT_PROMPT + "\n\nDocument content:\n" + text[:12000])
    elif ext == "doc":
        raise ValueError("Eski .doc format qo'llab-quvvatlanmaydi. Iltimos, .docx yoki PDF yuboring.")
    elif ext in ("xlsx", "xlsm"):
        try:
            text = _read_xlsx(file_bytes)
        except Exception as e:
            text = "[Could not parse: " + str(e) + "]"
        response = await generate_with_retry(model, EXTRACT_PROMPT + "\n\nSpreadsheet content:\n" + text[:12000])
    elif ext == "xls":
        raise ValueError("Eski .xls format qo'llab-quvvatlanmaydi. Iltimos, .xlsx, CSV yoki PDF yuboring.")
    elif ext == "csv":
        text = file_bytes.decode("utf-8", errors="replace")[:12000]
        response = await generate_with_retry(model, EXTRACT_PROMPT + "\n\nSpreadsheet content:\n" + text)
    else:
        text = file_bytes.decode("utf-8", errors="replace")[:12000]
        response = await generate_with_retry(model, EXTRACT_PROMPT + "\n\nFile content:\n" + text)

    return parse_json_response(response.text)


@require_auth
async def start(update, context):
    await update.message.reply_text(
        "👋 Dispatcher Bot ready!\n\n"
        "Send me a load confirmation (photo, PDF, Word, Excel) and I'll format it.\n"
        "I'll also ask your current location to calculate empty & loaded miles.\n\n"
        "/help for more info."
    )
    return ConversationHandler.END


@require_auth
async def help_cmd(update, context):
    await update.message.reply_text(
        "📌 Send any load/rate confirmation:\n"
        "🖼 Photo / screenshot\n"
        "📋 PDF\n"
        "📄 Word (.docx)\n"
        "📊 Excel (.xlsx) / CSV\n\n"
        "The bot will:\n"
        "1️⃣ Extract all load info\n"
        "2️⃣ Ask your current location\n"
        "3️⃣ Calculate empty & loaded miles\n"
        "4️⃣ Send formatted dispatcher message\n\n"
        "/cancel to cancel current operation"
    )
    return ConversationHandler.END


@require_auth
async def receive_file(update, context):
    msg = update.message
    caption = msg.caption or ""
    await msg.reply_text("⏳ Reading document...")
    try:
        if msg.document:
            f = await msg.document.get_file()
            file_bytes = bytes(await f.download_as_bytearray())
            filename = msg.document.file_name or "file"
        else:
            f = await msg.photo[-1].get_file()
            file_bytes = bytes(await f.download_as_bytearray())
            filename = "photo.jpg"

        load_data = await extract_load_data(file_bytes, filename, caption)
        context.user_data.clear()
        context.user_data["load_data"] = load_data

        pu_count = len(load_data.get("pickups", []))
        do_count = len(load_data.get("deliveries", []))
        await msg.reply_text(
            "✅ Found: " + str(pu_count) + " pickup(s), " + str(do_count) + " delivery stop(s)\n\n"
            "📍 What's your current location?\n"
            "(Type city name or full address, e.g. 'Dallas, TX')"
        )
        return WAITING_LOCATION
    except Exception as e:
        logger.exception("File processing error")
        await msg.reply_text("❌ Error reading file: " + str(e))
        return ConversationHandler.END


@require_auth
async def receive_location(update, context):
    current_location = update.message.text.strip()
    load_data = context.user_data.get("load_data", {})
    await update.message.reply_text("🗺 Calculating miles...")
    try:
        empty_miles, loaded_miles = await calculate_miles(current_location, load_data)
        final_message = build_message(load_data, empty_miles, loaded_miles)
        await update.message.reply_text(final_message, parse_mode="HTML")
        context.user_data["current_location"] = current_location
        context.user_data["empty_miles"] = empty_miles
        context.user_data["loaded_miles"] = loaded_miles
    except Exception as e:
        logger.exception("Miles calculation error")
        await update.message.reply_text("❌ Error: " + str(e))
        context.user_data.clear()
    return ConversationHandler.END


@require_auth
async def cancel(update, context):
    context.user_data.clear()
    await update.message.reply_text("❌ Cancelled. Send a new load document whenever you're ready.")
    return ConversationHandler.END


def _addresses_changed(old_data, new_data):
    def addr_set(data):
        addrs = []
        for pu in data.get("pickups", []):
            addrs.append((pu.get("address_line1", ""), pu.get("address_line2", "")))
        for do_ in data.get("deliveries", []):
            addrs.append((do_.get("address_line1", ""), do_.get("address_line2", "")))
        return addrs
    return addr_set(old_data) != addr_set(new_data)


async def edit_load_data(load_data, user_text):
    prompt = EDIT_PROMPT.format(
        load_data=json.dumps(load_data, ensure_ascii=False),
        user_text=user_text,
    )
    response = await generate_with_retry(model, prompt)
    parsed = parse_json_response(response.text)
    if isinstance(parsed, dict) and parsed.get("__not_an_edit__"):
        return None, False
    return parsed, True


@require_auth
async def chat_fallback(update, context):
    if not update.message or not update.message.text:
        return
    user_text = update.message.text.strip()
    if not user_text:
        return

    load_data = context.user_data.get("load_data")

    try:
        if load_data:
            updated_data, was_edit = await edit_load_data(load_data, user_text)
            if was_edit:
                context.user_data["load_data"] = updated_data

                if _addresses_changed(load_data, updated_data) and context.user_data.get("current_location"):
                    await update.message.reply_text("🗺 Address changed — recalculating miles...")
                    empty_miles, loaded_miles = await calculate_miles(
                        context.user_data["current_location"], updated_data
                    )
                    context.user_data["empty_miles"] = empty_miles
                    context.user_data["loaded_miles"] = loaded_miles
                else:
                    empty_miles = context.user_data.get("empty_miles", 0)
                    loaded_miles = context.user_data.get("loaded_miles", 0)

                final_message = build_message(updated_data, empty_miles, loaded_miles)
                await update.message.reply_text("✅ Updated:")
                await update.message.reply_text(final_message, parse_mode="HTML")
                return

            prompt = (
                "You are a helpful assistant for a freight dispatcher. "
                "The dispatcher is currently working on this load (JSON context):\n"
                + json.dumps(load_data, ensure_ascii=False)
                + "\n\nAnswer the dispatcher's question or message naturally and helpfully, "
                "using the load context above when relevant:\n\n" + user_text
            )
        else:
            prompt = (
                "You are a helpful, friendly assistant chatting with a freight dispatcher. "
                "Answer naturally and helpfully:\n\n" + user_text
            )

        response = await generate_with_retry(chat_model, prompt)
        reply = response.text.strip() if response.text else "🤔 I couldn't come up with a response to that."
        await update.message.reply_text(reply)
    except Exception as e:
        logger.exception("Chat fallback error")
        await update.message.reply_text("❌ Error: " + str(e))


def main():
    missing = [n for n, v in [
        ("TELEGRAM_TOKEN", TELEGRAM_TOKEN),
        ("GEMINI_KEY", GEMINI_KEY),
        ("GOOGLE_MAPS_KEY", GOOGLE_MAPS_KEY),
        ("BOT_PASSWORD", BOT_PASSWORD),
    ] if not v]
    if missing:
        logger.error("Missing Fly.io secrets: %s", ", ".join(missing))
    if not TELEGRAM_TOKEN:
        raise SystemExit("TELEGRAM_TOKEN secret is not set")

    app = Application.builder().token(TELEGRAM_TOKEN).build()
    conv_handler = ConversationHandler(
        entry_points=[
            MessageHandler(filters.Document.ALL, receive_file),
            MessageHandler(filters.PHOTO, receive_file),
        ],
        states={
            WAITING_LOCATION: [MessageHandler(filters.TEXT & ~filters.COMMAND, receive_location)]
        },
        fallbacks=[
            CommandHandler("cancel", cancel),
            CommandHandler("start", start),
        ],
    )
    app.add_handler(CommandHandler("start", start))
    app.add_handler(CommandHandler("help", help_cmd))
    app.add_handler(conv_handler)
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, chat_fallback))
    logger.info("✅ Bot running…")
    app.run_polling(allowed_updates=Update.ALL_TYPES, drop_pending_updates=True)


if __name__ == "__main__":
    main()
